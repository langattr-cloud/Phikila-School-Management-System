"""Parse KCB statement uploads for preview before reconciliation.

Supports CSV, XLSX/XLSM, and text-based PDF statements. Scanned/image-only PDFs
are rejected with a clear message rather than silently importing guessed data.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.modules.scheduling.tenancy import Principal, require_role

router = APIRouter()
MAX_UPLOAD_BYTES = 15 * 1024 * 1024


class ParsedTransaction(BaseModel):
    reference: str
    amount: Decimal
    transaction_date: datetime | None = None


class StatementPreview(BaseModel):
    filename: str
    file_type: str
    parsed_rows: int
    transactions: list[ParsedTransaction]
    warnings: list[str] = Field(default_factory=list)


def _norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().casefold()).strip()


def _money(value: object) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    raw = str(value).strip().replace(",", "").replace("KES", "").replace("Ksh", "").replace("ksh", "").strip()
    if raw.startswith("(") and raw.endswith(")"):
        raw = "-" + raw[1:-1]
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError):
        return None


def _date(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y %H:%M", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return None


def _reference(value: object) -> str | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    # Keep a whole transaction reference when the cell contains just the reference.
    if re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9./_-]{5,49}", raw):
        return raw
    # Prefer the explicit payment reference commonly included in KCB M-PESA narrations.
    explicit = re.search(r"(?:M\s*[- ]?PESA\s*)?(?:REF(?:ERENCE)?|TRANS(?:ACTION)?\s*(?:ID|REF))\s*[:#-]?\s*([A-Z0-9]{8,14})", raw, re.IGNORECASE)
    if explicit:
        return explicit.group(1).upper()
    # Otherwise accept a single reference-like token, not the first word in a narration.
    matches = re.findall(r"(?<![A-Za-z0-9])([A-Z0-9]{8,14})(?![A-Za-z0-9])", raw.upper())
    return matches[-1] if matches else None


def _rows_to_transactions(rows: list[list[object]]) -> tuple[list[ParsedTransaction], list[str]]:
    warnings: list[str] = []
    if not rows:
        raise HTTPException(422, "No transaction rows were found in the statement.")
    header_index = None
    header = []
    for index, row in enumerate(rows[:40]):
        normalized = [_norm(x) for x in row]
        has_ref = any(any(term in cell for term in ("reference", "ref no", "transaction id", "transaction code", "narration", "description", "details")) for cell in normalized)
        has_money = any(any(term in cell for term in ("credit", "deposit", "money in", "amount", "paid in")) for cell in normalized)
        if has_ref and has_money:
            header_index, header = index, normalized
            break
    if header_index is None:
        raise HTTPException(422, "Could not identify statement columns. Use a statement with headings for transaction reference/narration and credit/amount.")

    def find_col(terms: tuple[str, ...]) -> int | None:
        for idx, cell in enumerate(header):
            if any(term == cell or term in cell for term in terms):
                return idx
        return None

    ref_col = find_col(("transaction reference", "reference number", "reference", "ref no", "transaction id", "transaction code", "narration", "description", "details", "remarks"))
    credit_col = find_col(("credit amount", "credit", "deposit", "money in", "paid in", "amount credited"))
    amount_col = find_col(("transaction amount", "amount"))
    debit_col = find_col(("debit amount", "debit", "withdrawal", "money out", "paid out"))
    date_col = find_col(("transaction date", "date", "value date", "posting date", "time"))
    type_col = find_col(("transaction type", "type", "dr cr", "credit debit"))

    if ref_col is None or (credit_col is None and amount_col is None):
        raise HTTPException(422, "The statement must include a reference/narration column and a credit or amount column.")
    if credit_col is None and type_col is None and debit_col is None:
        raise HTTPException(422, "This statement has an amount column but does not identify credits versus debits. Export a version with a Credit column or transaction type.")

    transactions: list[ParsedTransaction] = []
    skipped = 0
    for row in rows[header_index + 1:]:
        if not row or not any(str(v or "").strip() for v in row):
            continue
        ref = _reference(row[ref_col] if ref_col < len(row) else None)
        credit = _money(row[credit_col] if credit_col is not None and credit_col < len(row) else None)
        amount = credit
        if amount is None and amount_col is not None:
            amount = _money(row[amount_col] if amount_col < len(row) else None)
            tx_type = _norm(row[type_col] if type_col is not None and type_col < len(row) else "")
            if type_col is not None and credit_col is None:
                if not any(token in tx_type for token in ("credit", "deposit", "receipt", " cr")):
                    amount = None
            elif debit_col is not None and credit_col is None:
                debit = _money(row[debit_col] if debit_col < len(row) else None)
                if debit is not None and amount == debit:
                    amount = None
        if not ref or amount is None or amount <= 0:
            skipped += 1
            continue
        tx_date = _date(row[date_col] if date_col is not None and date_col < len(row) else None)
        transactions.append(ParsedTransaction(reference=ref, amount=amount, transaction_date=tx_date))
    if not transactions:
        raise HTTPException(422, "No credit transactions could be parsed. Confirm the file has a reference and a positive credit amount for each row.")
    if skipped:
        warnings.append(f"{skipped} row(s) were skipped because they did not look like a positive credit with a transaction reference.")
    seen: set[str] = set()
    unique: list[ParsedTransaction] = []
    for transaction in transactions:
        if transaction.reference in seen:
            raise HTTPException(422, f"Duplicate reference {transaction.reference} appears in the uploaded statement. Check the file before reconciling.")
        seen.add(transaction.reference)
        unique.append(transaction)
    return unique, warnings


def _parse_csv(data: bytes) -> list[list[object]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _parse_excel(data: bytes) -> list[list[object]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    except Exception as exc:
        raise HTTPException(422, "Could not read this Excel workbook. Save it as .xlsx and try again.") from exc


def _parse_pdf(data: bytes) -> list[list[object]]:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            rows: list[list[object]] = []
            for page in pdf.pages:
                for table in (page.extract_tables() or []):
                    rows.extend([list(row or []) for row in table])
            if rows:
                return rows
            # Text fallback preserves the visible columns on many digital statements.
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        raise HTTPException(422, "Could not read this PDF. Try exporting the statement again from KCB.") from exc

    rows = []
    for line in text.splitlines():
        cells = [cell.strip() for cell in re.split(r"\s{2,}|\t+", line.strip()) if cell.strip()]
        if len(cells) >= 2:
            rows.append(cells)
    if not rows:
        raise HTTPException(422, "This PDF appears to be scanned or image-only. Upload a text-based PDF or export it to CSV/Excel.")
    return rows


@router.post("/finance/payment-inbox/reconcile-upload/preview", response_model=StatementPreview)
async def preview_kcb_statement_upload(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_role("admin")),
):
    """Parse an uploaded statement and return a preview; does not change any records."""
    del db, principal  # Auth and tenant scope are required even though preview is read-only.
    filename = Path(file.filename or "statement").name
    suffix = Path(filename).suffix.casefold()
    if suffix not in {".csv", ".xlsx", ".xlsm", ".pdf"}:
        raise HTTPException(415, "Supported statement formats are CSV, Excel (.xlsx/.xlsm), and PDF.")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Statement file exceeds the 15 MB upload limit.")
    if not data:
        raise HTTPException(422, "The uploaded statement file is empty.")
    if suffix == ".csv":
        rows = _parse_csv(data)
        file_type = "CSV"
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _parse_excel(data)
        file_type = "Excel"
    else:
        rows = _parse_pdf(data)
        file_type = "PDF"
    transactions, warnings = _rows_to_transactions(rows)
    return StatementPreview(filename=filename, file_type=file_type, parsed_rows=len(transactions), transactions=transactions, warnings=warnings)
