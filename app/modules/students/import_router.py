"""Bulk student import from Excel workbooks."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
import re

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.modules.scheduling.tenancy import Principal, require_role
from app.modules.academics.models import AcademicYear, Level, SchoolClass, Term
from . import models_v2 as m

router = APIRouter()

MAX_FILE_BYTES = 10 * 1024 * 1024
REQUIRED_HEADERS = {"admission_number", "first_name", "last_name", "academic_year", "level", "class"}
HEADER_ALIASES = {
    "admissionnumber": "admission_number", "admissionno": "admission_number", "admno": "admission_number",
    "admission_number": "admission_number", "first_name": "first_name", "firstname": "first_name",
    "middle_name": "middle_name", "middlename": "middle_name", "last_name": "last_name", "lastname": "last_name",
    "preferred_name": "preferred_name", "preferredname": "preferred_name", "dob": "date_of_birth",
    "dateofbirth": "date_of_birth", "date_of_birth": "date_of_birth", "gender": "gender", "email": "email",
    "phone": "phone", "address": "address", "nationality": "nationality", "nationalid": "national_id",
    "national_id": "national_id", "admissiondate": "admission_date", "admission_date": "admission_date",
    "status": "status", "academic_year": "academic_year", "academicyear": "academic_year", "academic_year_id": "academic_year_id",
    "level": "level", "level_name": "level", "level_code": "level_code", "level_id": "level_id",
    "class": "class", "class_name": "class", "class_code": "class_code", "class_id": "class_id",
    "term": "term", "term_name": "term", "term_id": "term_id",
}


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _excel_date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return date(1899, 12, 30) + timedelta(days=float(value))
    raw = _text(value).replace(".", "/").replace("-", "/")
    for fmt in ("%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Invalid date '{raw}'. Use DD/MM/YYYY or YYYY-MM-DD.")


def _lookup_id(value: str, by_id: dict[int, object]) -> int | None:
    if not value:
        return None
    try:
        candidate = int(float(value))
    except ValueError:
        return None
    return candidate if candidate in by_id else None


def _find_year(db: Session, school_id: int, row: dict[str, str]) -> AcademicYear | None:
    if row.get("academic_year_id"):
        found_id = _lookup_id(row["academic_year_id"], {x.id: x for x in db.query(AcademicYear).filter(AcademicYear.school_id == school_id).all()})
        if found_id:
            return db.get(AcademicYear, found_id)
    value = row.get("academic_year", "")
    if not value:
        return None
    return db.query(AcademicYear).filter(AcademicYear.school_id == school_id, func.lower(AcademicYear.name) == value.lower()).first()


def _find_level(db: Session, school_id: int, row: dict[str, str]) -> Level | None:
    levels = db.query(Level).filter(Level.school_id == school_id).all()
    if row.get("level_id"):
        found_id = _lookup_id(row["level_id"], {x.id: x for x in levels})
        if found_id:
            return db.get(Level, found_id)
    value = row.get("level", "")
    if not value:
        return None
    value_l = value.lower()
    return next((x for x in levels if x.name.lower() == value_l or x.code.lower() == value_l), None)


def _find_class(db: Session, school_id: int, year_id: int, level_id: int, row: dict[str, str]) -> SchoolClass | None:
    classes = db.query(SchoolClass).filter(
        SchoolClass.school_id == school_id,
        SchoolClass.academic_year_id == year_id,
        SchoolClass.level_id == level_id,
    ).all()
    if row.get("class_id"):
        found_id = _lookup_id(row["class_id"], {x.id: x for x in classes})
        if found_id:
            return db.get(SchoolClass, found_id)
    code = row.get("class_code", "").lower()
    name = row.get("class", "").lower()
    return next((x for x in classes if (code and x.code.lower() == code) or (name and x.name.lower() == name)), None)


def _find_term(db: Session, school_id: int, year_id: int, row: dict[str, str]) -> Term | None:
    terms = db.query(Term).filter(Term.school_id == school_id, Term.academic_year_id == year_id).all()
    if row.get("term_id"):
        found_id = _lookup_id(row["term_id"], {x.id: x for x in terms})
        if found_id:
            return db.get(Term, found_id)
    value = row.get("term", "").lower()
    if value:
        return next((x for x in terms if x.name.lower() == value), None)
    return next((x for x in terms if x.is_current), None)


@router.post("/students/import", status_code=status.HTTP_200_OK)
async def import_students_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_role("admin", "scheduler")),
):
    """Validate and import an .xlsx workbook atomically for the current school."""
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Please upload an Excel .xlsx file.")
    content = await file.read()
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "The Excel file must be 10 MB or smaller.")
    if not content:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The Excel file is empty.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if not raw_headers:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "The first worksheet has no header row.")
        headers: list[str] = []
        for header in raw_headers:
            normalized = HEADER_ALIASES.get(_key(header), _key(header))
            headers.append(normalized)
        missing = sorted(REQUIRED_HEADERS - set(headers))
        if missing:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"Missing required columns: {', '.join(missing)}.")

        parsed: list[tuple[int, dict[str, str]]] = []
        for row_number, values in enumerate(rows, start=2):
            if not any(v not in (None, "") for v in values):
                continue
            row = {headers[i]: _text(values[i]) for i in range(min(len(headers), len(values))) if headers[i]}
            row["__row"] = str(row_number)
            parsed.append((row_number, row))
        if not parsed:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "No student rows were found below the header row.")
        if len(parsed) > 2000:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "A single import can contain at most 2,000 students.")

        errors: list[str] = []
        admission_numbers: set[str] = set()
        prepared: list[dict[str, object]] = []
        for row_number, row in parsed:
            admission = row.get("admission_number", "").strip()
            first = row.get("first_name", "").strip()
            last = row.get("last_name", "").strip()
            if not admission or not first or not last:
                errors.append(f"Row {row_number}: admission number, first name and last name are required.")
                continue
            admission_key = admission.lower()
            if admission_key in admission_numbers:
                errors.append(f"Row {row_number}: duplicate admission number '{admission}'.")
                continue
            admission_numbers.add(admission_key)
            existing = db.query(m.Student).filter(m.Student.school_id == principal.school_id, func.lower(m.Student.admission_number) == admission_key).first()
            if existing:
                errors.append(f"Row {row_number}: admission number '{admission}' already exists.")
                continue
            year = _find_year(db, principal.school_id, row)
            level = _find_level(db, principal.school_id, row)
            school_class = _find_class(db, principal.school_id, year.id, level.id, row) if year and level else None
            term = _find_term(db, principal.school_id, year.id, row) if year else None
            if not year:
                errors.append(f"Row {row_number}: academic year '{row.get('academic_year', '')}' was not found.")
                continue
            if not level:
                errors.append(f"Row {row_number}: level '{row.get('level', '')}' was not found.")
                continue
            if not school_class:
                errors.append(f"Row {row_number}: class '{row.get('class', '') or row.get('class_code', '')}' was not found for that academic year and level.")
                continue
            try:
                dob = _excel_date(row.get("date_of_birth"))
                admission_date = _excel_date(row.get("admission_date")) or date.today()
            except ValueError as exc:
                errors.append(f"Row {row_number}: {exc}")
                continue
            prepared.append({"row_number": row_number, "row": row, "year": year, "level": level, "class": school_class, "term": term, "dob": dob, "admission_date": admission_date})

        if errors:
            workbook.close()
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, {"message": "Import validation failed. No students were imported.", "errors": errors})

        for item in prepared:
            row = item["row"]
            student = m.Student(
                school_id=principal.school_id,
                admission_number=row["admission_number"],
                first_name=row["first_name"],
                middle_name=row.get("middle_name") or None,
                last_name=row["last_name"],
                preferred_name=row.get("preferred_name") or None,
                date_of_birth=item["dob"],
                gender=row.get("gender") or None,
                email=row.get("email") or None,
                phone=row.get("phone") or None,
                address=row.get("address") or None,
                nationality=row.get("nationality") or "Kenyan",
                national_id=row.get("national_id") or None,
                admission_date=item["admission_date"],
                status=(row.get("status") or "active").lower(),
            )
            db.add(student)
            db.flush()
            db.add(m.StudentEnrollment(
                school_id=principal.school_id,
                student_id=student.id,
                academic_year_id=item["year"].id,
                term_id=item["term"].id if item["term"] else None,
                level_id=item["level"].id,
                school_class_id=item["class"].id,
                class_id=item["class"].id,
                status="active",
                enrollment_date=item["admission_date"],
            ))
            from app.modules.scheduling.models import TtAuditEntry
            db.add(TtAuditEntry(
                school_id=principal.school_id,
                actor=principal.email or principal.user_id,
                action="create",
                entity="student",
                entity_id=student.id,
                summary=f"Imported student {student.first_name} {student.last_name} ({student.admission_number}) from Excel",
            ))
        db.commit()
        workbook.close()
        return {"imported": len(prepared), "rows": len(parsed), "message": f"Successfully imported {len(prepared)} students."}
    except HTTPException:
        raise
    except Exception as exc:
        db.rollback()
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"Could not read the Excel workbook: {exc}") from exc
